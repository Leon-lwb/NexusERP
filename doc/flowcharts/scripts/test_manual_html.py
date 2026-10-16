import gzip
import hashlib
import tempfile
import unittest
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document

try:
    from .gen_manual_html import generate_manual_html
except ImportError:
    from gen_manual_html import generate_manual_html


ROOT = Path(__file__).resolve().parents[3]
SOURCE = ROOT / "doc/system/user-manual/NexusERP_用户操作手册.docx"


class ManualHtmlGenerationTest(unittest.TestCase):

    def test_generates_complete_static_manual(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "manual"
            document = Document(SOURCE)
            result = generate_manual_html(SOURCE, output)

            index = output / "index.html"
            compressed = output / "index.html.gz"
            manifest = output / "manifest.json"
            self.assertTrue(index.is_file())
            self.assertTrue(compressed.is_file())
            self.assertTrue(manifest.is_file())
            self.assertEqual(gzip.decompress(compressed.read_bytes()), index.read_bytes())

            index_text = index.read_text(encoding="utf-8")
            soup = BeautifulSoup(index_text, "html.parser")
            self.assertEqual(soup.html.get("lang"), "zh-CN")
            self.assertEqual(soup.title.string, document.core_properties.title)
            self.assertTrue(soup.select_one('link[rel="icon"]')["href"].startswith("data:image/svg+xml,"))
            self.assertEqual(len(soup.select("article table")), len(document.tables))
            self.assertEqual(len(soup.select("article figure img")), 29)
            self.assertEqual(len(soup.select('article h1[data-kind="chapter"]')), 23)
            self.assertGreaterEqual(len(soup.select("nav a[href^='#']")), 200)
            self.assertIsNotNone(soup.select_one("#manual-search"))
            self.assertIsNotNone(soup.select_one("#diagram-dialog"))
            self.assertNotIn("目录 3", soup.get_text(" ", strip=True))
            self.assertRegex(index_text, r"article\s*\{[^}]*overflow-x:\s*clip")

            image_sources = [image["src"] for image in soup.select("article figure img")]
            self.assertEqual(len(image_sources), len(set(image_sources)))
            for source in image_sources:
                self.assertFalse(source.startswith(("http:", "https:", "file:")))
                self.assertTrue((output / source).is_file(), source)

            self.assertEqual(result["tables"], len(document.tables))
            self.assertEqual(result["images"], 29)
            self.assertEqual(result["version"], document.core_properties.version)
            self.assertIn("自研报表", soup.get_text())
            self.assertIn("工程设备与本地连接器", soup.get_text())
            self.assertEqual(result["chapters"], 23)
            self.assertEqual(
                result["source_sha256"],
                hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
            )


if __name__ == "__main__":
    unittest.main()
