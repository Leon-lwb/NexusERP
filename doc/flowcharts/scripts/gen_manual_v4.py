#!/usr/bin/env python3
"""Build the current manual with chapter-local, embedded process diagrams."""

import argparse
import re
from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

import gen_manual_v3 as base
from gen_current_flowcharts import DATE, ROOT, OUT as CHARTS, load_specs
from manual_current_warehouse import ch23_warehouse

OUTPUT = ROOT / 'doc/system/user-manual/NexusERP_用户操作手册.docx'
HEADING_1, HEADING_2 = base.h1, base.h2


def field(paragraph, instruction, placeholder=''):
    element = OxmlElement('w:fldSimple')
    element.set(qn('w:instr'), instruction)
    if placeholder:
        run = OxmlElement('w:r')
        text = OxmlElement('w:t')
        text.text = placeholder
        run.append(text)
        element.append(run)
    paragraph._p.append(element)


def configure(doc):
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    section.header_distance = section.footer_distance = Cm(0.9)
    section.different_first_page_header_footer = True
    for name in ('Normal', 'List Bullet', 'List Number', 'Caption'):
        style = doc.styles[name]
        style.font.name = base.FONT
        style.font.size = Pt(10.5 if name != 'Caption' else 9)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), base.FONT)
        style.paragraph_format.space_after = Pt(4)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.15
    for level in (1, 2, 3):
        style = doc.styles[f'Heading {level}']
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.space_before = Pt(10)
    header = section.header.paragraphs[0]
    base.set_font(header.add_run('NexusERP 用户操作手册 | V4.1'), size=8, color=RGBColor.from_string('62717A'))
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    base.set_font(footer.add_run('NexusERP  /  '), size=8)
    field(footer, ' PAGE ', '1')
    settings = doc.settings.element
    update = OxmlElement('w:updateFields')
    update.set(qn('w:val'), 'true')
    settings.append(update)
    doc.core_properties.title = 'NexusERP 用户操作手册 V4.1'
    doc.core_properties.subject = 'Current operations and evidence-backed process diagrams'
    doc.core_properties.author = 'NexusERP'
    doc.core_properties.version = '4.1'
    doc.core_properties.comments = 'Reports/engineering baseline: plus/b154279b, 2026-10-17; other chapters/diagrams: plus/fe59e4ae, 2026-10-12. Not production acceptance.'


class Chapters:
    def __init__(self, specs):
        self.by_section = defaultdict(list)
        self.used = set()
        self.specs = specs
        for spec in specs:
            self.by_section[spec['chapter']].append(spec)

    def h1(self, doc, text):
        needs_break = bool(doc.paragraphs) and not doc.paragraphs[-1]._p.xpath('.//w:br[@w:type="page"]')
        HEADING_1(doc, text)
        if needs_break:
            doc.paragraphs[-1].paragraph_format.page_break_before = True

    def h2(self, doc, text):
        match = re.match(r'^(\d+\.\d+)\b', text)
        specs = self.by_section.get(match.group(1), []) if match else []
        HEADING_2(doc, text)
        if specs:
            doc.paragraphs[-1].paragraph_format.page_break_before = True
        for index, spec in enumerate(specs):
            if index:
                doc.add_page_break()
            path = CHARTS / Path(spec['path']).with_suffix('.preview.png')
            if not path.is_file():
                raise FileNotFoundError(f'Render diagrams first: {path}')
            with Image.open(path) as image:
                width, height = image.size
            image_width = min(Cm(16.6), int(Cm(21.7) * width / height))
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.space_after = Pt(4)
            shape = paragraph.add_run().add_picture(str(path), width=image_width)
            shape._inline.docPr.set('descr', f"{spec['title']}。{spec['summary']}")
            caption = doc.add_paragraph(style='Caption')
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            number = next(i for i, item in enumerate(self.specs, 1) if item['id'] == spec['id'])
            base.set_font(caption.add_run(f"图 {number:02d}  {spec['title']}"), size=9)
            self.used.add(spec['id'])
        if specs:
            doc.add_page_break()
            for spec in specs:
                base.para(doc, '流程说明：' + spec['summary'])
                if spec.get('prerequisites'):
                    base.para(doc, '使用前：' + '；'.join(spec['prerequisites']), size=9.5)
                if spec.get('caveats'):
                    base.para(doc, '使用边界：' + '；'.join(spec['caveats']), size=9.5)

    def args(self, doc):
        return (doc, self.h1, self.h2, base.h3, base.para, base.bullet, base.numbered, base.table)


def front_matter(doc):
    for _ in range(4):
        doc.add_paragraph()
    base.para(doc, 'NexusERP', size=30, bold=True, color=base.ACCENT, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
    base.para(doc, '用户操作手册', size=24, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
    base.para(doc, 'ERP · MES · IoT · AI', size=14, color=RGBColor.from_string('62717A'), align=WD_ALIGN_PARAGRAPH.CENTER, after=34)
    for text in ('文档版本：V4.1', '报表/工程核对日期：2026-10-17', '报表/工程代码基线：plus / b154279b', '编制：项目文档维护'):
        base.para(doc, text, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    base.para(doc, '当前代码的操作指南，不替代现场部署与人工验收。', size=10, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
    doc.add_page_break()
    base.h1(doc, '阅读与使用边界')
    base.para(doc, '面向业务操作人员、管理员和现场支持人员。原有二十章保留，追加工位/履历、签章和仓储章节；'
                   '各流程图放置在关联章节，并保留可编辑源图。角色授权和环境菜单不同，实际可见入口可能不同。')
    base.table(doc, ['事项', '本手册口径'], [
        ['代码事实', '自研报表与工程连接器：plus/b154279b，2026-10-17；其他章节与29张流程图沿用2026-10-12的fe59e4ae基线。关联源码见图集来源索引。'],
        ['普通用户', '依据按钮、实际资料和当前对象状态执行；看见菜单或查询结果不等于可以写入。'],
        ['配置 / 部署', '模型、Vision、MQ、OSS、设备、SnailJob 和流程定义需管理员另行配置与核对。'],
        ['人工验收', '没有验证生产数据、Station 真机、移动端全链路或签章人工端到端验收。'],
        ['安全操作', '不提供通用演示账户、密钥或迁移命令；存量库不重导初始化 SQL。'],
        ['图形含义', '蓝色为人工操作，绿色为系统处理，黄色为校验/审批，灰色为状态/结果；虚线为异步或异常分支。'],
    ], widths=[3.2, 12.8])
    doc.add_page_break()
    base.h1(doc, '目录')
    doc.paragraphs[-1].style = doc.styles['Title']
    field(doc.add_paragraph(), ' TOC \\o "1-2" \\h \\z \\u ', '请在 Word/WPS 中更新整个目录。')
    doc.add_page_break()


def diagram_appendix(doc, chapters):
    chapters.h1(doc, '附录C 流程图索引与使用条件')
    base.para(doc, '全部流程图已嵌入本手册，无需联网加载图片。可编辑 .drawio 与对应 .preview.png 位于 doc/flowcharts。'
                   '源码证据和逐图条件参见同目录的 来源与使用边界.md；源码存在不是生产启用证据。')
    base.table(doc, ['图号', '关联章节', '流程图'], [
        [f'{i:02d}', spec['chapter'], spec['title']] for i, spec in enumerate(chapters.specs, 1)
    ], widths=[1.3, 2.3, 12.4])
    base.para(doc, '维护入口：doc/flowcharts/scripts/gen_current_flowcharts.py 和 gen_manual_v4.py。'
                   '修改时先核对真实页面、服务和菜单，再更新 current_*.json 和对应章节，不能只改预览图。')


def normalize_tables(doc):
    section = doc.sections[0]
    available = section.page_width - section.left_margin - section.right_margin
    for table in doc.tables:
        requested = [cell.width or column.width for cell, column in zip(table.rows[0].cells, table.columns)]
        total = sum(requested)
        widths = [int(available * width / total) for width in requested]
        table.autofit = False
        for i, column in enumerate(table.columns):
            column.width = widths[i]
        for row in table.rows:
            for i, cell in enumerate(row.cells):
                cell.width = widths[i]
        table_width = table._tbl.tblPr.find(qn('w:tblW'))
        table_width.set(qn('w:type'), 'dxa')
        table_width.set(qn('w:w'), str(round(available / 635)))
        header = table.rows[0]._tr.get_or_add_trPr()
        repeat = OxmlElement('w:tblHeader')
        repeat.set(qn('w:val'), 'true')
        header.append(repeat)
        for row in table.rows:
            properties = row._tr.get_or_add_trPr()
            properties.append(OxmlElement('w:cantSplit'))
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(3)


def normalize_numbering(doc):
    numbering = doc.part.numbering_part.element
    style_id = doc.styles['List Number']._element.xpath('.//w:numId')[0].val
    source = next(node for node in numbering if node.tag == qn('w:num') and int(node.get(qn('w:numId'))) == style_id)
    abstract_id = source.find(qn('w:abstractNumId')).get(qn('w:val'))
    next_id = max(int(node.get(qn('w:numId'))) for node in numbering if node.tag == qn('w:num')) + 1
    active = None
    for paragraph in doc.paragraphs:
        if paragraph.style.name != 'List Number':
            active = None
            continue
        if active is None:
            active = next_id
            next_id += 1
            num = OxmlElement('w:num')
            num.set(qn('w:numId'), str(active))
            abstract = OxmlElement('w:abstractNumId')
            abstract.set(qn('w:val'), abstract_id)
            num.append(abstract)
            override = OxmlElement('w:lvlOverride')
            override.set(qn('w:ilvl'), '0')
            start = OxmlElement('w:startOverride')
            start.set(qn('w:val'), '1')
            override.append(start)
            num.append(override)
            numbering.append(num)
        properties = paragraph._p.get_or_add_pPr().get_or_add_numPr()
        properties.get_or_add_ilvl().val = 0
        properties.get_or_add_numId().val = active


def build(output=OUTPUT):
    from manual_current_operations import ch21_station, extend_iot, extend_cv
    from manual_current_intelligence import extend_ai, ch22_esign
    from manual_current_engineering import extend_engineering
    from manual_current_report import ch19_report
    specs = load_specs()
    if len(specs) < 29:
        raise ValueError('Current diagram set is incomplete')
    doc = Document()
    configure(doc)
    front_matter(doc)
    chapters = Chapters(specs)
    original_h1, original_h2 = base.h1, base.h2
    try:
        # The overview calls its helpers directly rather than receiving callbacks.
        base.h1, base.h2 = chapters.h1, chapters.h2
        base.ch1_overview(doc)
    finally:
        base.h1, base.h2 = original_h1, original_h2
    callbacks = chapters.args(doc)
    for chapter in (base.ch2_bd, base.ch3_so, base.ch4_pu, base.ch5_req, base.ch6_prd,
                    base.ch7_sch, base.ch8_mes, base.ch9_qc, base.ch10_eq, base.ch11_lt):
        chapter(*callbacks)
    base.ch12_iot(*callbacks)
    extend_iot(*callbacks)
    base.ch13_cv(*callbacks)
    extend_cv(*callbacks)
    base.ch14_ai(*callbacks)
    extend_ai(*callbacks)
    extend_engineering(*callbacks)
    for chapter in (base.ch15_im, base.ch16_fin, base.ch17_pay, base.ch18_workflow,
                    ch19_report, base.ch20_system, ch21_station, ch22_esign, ch23_warehouse):
        chapter(*callbacks)
    base.appendix(*callbacks)
    diagram_appendix(doc, chapters)
    missing = {s['id'] for s in specs} - chapters.used
    if missing:
        raise ValueError(f'Diagrams not placed in a real manual section: {sorted(missing)}')
    normalize_tables(doc)
    normalize_numbering(doc)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(f'Built {output}: {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables, {len(doc.inline_shapes)} embedded diagrams')
    return doc


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--output', type=Path, default=OUTPUT)
    args = parser.parse_args()
    build(args.output)


if __name__ == '__main__':
    main()
